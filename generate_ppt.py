import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    NAVY_DARK = RGBColor(15, 23, 42)       # #0F172A
    NAVY_LIGHT = RGBColor(30, 41, 59)      # #1E293B
    SAFFRON = RGBColor(230, 81, 0)         # #E65100
    GOLD = RGBColor(245, 158, 11)          # #F59E0B
    EMERALD = RGBColor(16, 185, 129)       # #10B981
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(248, 250, 252)     # #F8FAFC
    CARD_BORDER = RGBColor(226, 232, 240)  # #E2E8F0
    TEXT_DARK = RGBColor(15, 23, 42)
    TEXT_MUTED = RGBColor(100, 116, 139)   # #64748B
    BLUE_ACCENT = RGBColor(37, 99, 235)    # #2563EB

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title_text, category_text, dark=False):
        # Header category badge
        badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_b = badge_box.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = tf_b.margin_right = tf_b.margin_top = tf_b.margin_bottom = 0
        p_b = tf_b.paragraphs[0]
        p_b.text = category_text.upper()
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.font.color.rgb = GOLD if dark else SAFFRON

        # Main slide title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.7))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE if dark else NAVY_DARK

    def add_card(slide, left, top, width, height, bg_color=WHITE, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        return shape

    # =========================================================================
    # SLIDE 1: Title, Team Details & Problem Statement Overview
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, NAVY_DARK)

    # Top Pill Badge
    pill = add_card(s1, Inches(0.8), Inches(0.6), Inches(4.8), Inches(0.45), bg_color=NAVY_LIGHT, border_color=SAFFRON)
    tf = pill.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "★ PROBLEM STATEMENT: YAS-01 | MYAS & SAI ALIGNED"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = GOLD

    # Main Project Title
    tbox = s1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(1.3))
    tf = tbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "AI4Seva: AI-Based Sports Talent Identification Portal"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = WHITE

    p2 = tf.add_paragraph()
    p2.text = "Decentralized Grassroots Athlete Scouting & Biomechanical Kinematics Engine"
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(148, 163, 184)

    # Team Info Strip Card
    strip = add_card(s1, Inches(0.8), Inches(2.55), Inches(11.73), Inches(0.65), bg_color=NAVY_LIGHT, border_color=BLUE_ACCENT)
    tf_strip = strip.text_frame
    tf_strip.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_strip = tf_strip.paragraphs[0]
    p_strip.text = "Team Number: [TEAM-ID]   |   Team Name: [TEAM-NAME]   |   Track: AI4Seva / Sports Authority of India (SAI)"
    p_strip.font.size = Pt(12)
    p_strip.font.bold = True
    p_strip.font.color.rgb = WHITE
    p_strip.alignment = PP_ALIGN.CENTER

    # 3 Problem Context Cards
    card_w = Inches(3.75)
    card_h = Inches(3.5)
    y_pos = Inches(3.4)

    # Card 1: Citizen/Athlete Side Problem
    c1 = add_card(s1, Inches(0.8), y_pos, card_w, card_h, bg_color=NAVY_LIGHT, border_color=RGBColor(51, 65, 85))
    tf1 = c1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_right = tf1.margin_top = Inches(0.25)
    p = tf1.paragraphs[0]
    p.text = "1. Grassroots Invisibility"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = GOLD
    
    bullets1 = [
        "Geographic Isolation: Talents in remote hill & tribal regions (e.g. Manipur districts like Ukhrul, Churachandpur) lack physical scouting camps.",
        "Prohibitive Costs: Rural families cannot afford travel & lodging for distant state-level trials.",
        "Delayed Identification: Potential champions discovered at ages 18-21 instead of peak developmental age (12-16)."
    ]
    for b in bullets1:
        p_b = tf1.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = RGBColor(226, 232, 240)
        p_b.space_before = Pt(8)

    # Card 2: Conventional Selection Gaps
    c2 = add_card(s1, Inches(4.79), y_pos, card_w, card_h, bg_color=NAVY_LIGHT, border_color=RGBColor(51, 65, 85))
    tf2 = c2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_right = tf2.margin_top = Inches(0.25)
    p = tf2.paragraphs[0]
    p.text = "2. Conventional Selection Gaps"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = SAFFRON
    
    bullets2 = [
        "Subjective Bias & Nepotism: Scouting relies heavily on manual coach intuition and unstandardized visual impressions.",
        "Fragmented Records: Paper forms (SAI Form K-01) and local club match data are lost across isolated federations.",
        "No Longitudinal Tracking: Inability to monitor physiological progress, injury patterns, or developmental trajectories."
    ]
    for b in bullets2:
        p_b = tf2.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = RGBColor(226, 232, 240)
        p_b.space_before = Pt(8)

    # Card 3: AI4Seva Proposed Solution
    c3 = add_card(s1, Inches(8.78), y_pos, card_w, card_h, bg_color=NAVY_LIGHT, border_color=EMERALD)
    tf3 = c3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = tf3.margin_right = tf3.margin_top = Inches(0.25)
    p = tf3.paragraphs[0]
    p.text = "3. AI4Seva Value Proposition"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    
    bullets3 = [
        "Ubiquitous Mobile Assessment: Any smartphone records standard fitness battery drills & video kinematics.",
        "Objective AI Talent Index: Calculates multi-factor scores, peer percentiles, and multi-sport suitability.",
        "Unified National Khel ID: Standardized athlete registry feeding high-potential prospects directly to qualified SAI coaches."
    ]
    for b in bullets3:
        p_b = tf3.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = RGBColor(226, 232, 240)
        p_b.space_before = Pt(8)


    # =========================================================================
    # SLIDE 2: Solution Architecture & End-to-End Workflow Diagram
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, LIGHT_BG)
    add_header(s2, "End-to-End System Architecture & Scouting Workflow", "TECHNICAL ARCHITECTURE & WORKFLOW")

    # 5 Sequential Workflow Step Cards
    step_w = Inches(2.2)
    step_h = Inches(4.3)
    y_step = Inches(1.6)
    gap = Inches(0.18)

    steps_data = [
        ("Step 1: Onboard & ID", SAFFRON, [
            "Athlete/School PE registration",
            "Generate unique Khel ID (e.g. KHEL-IND-MN-2026-XXXX)",
            "Aadhaar-linked profile & district mapping",
            "Upload birth cert & club history"
        ]),
        ("Step 2: Dual Ingestion", BLUE_ACCENT, [
            "Standard Fitness Battery (SAI Form K-01: Sprint, Jump, Beep Test, Wingspan)",
            "Mobile Video Capture of physical exercise & technique",
            "Match records & tournament stats"
        ]),
        ("Step 3: AI Engine", NAVY_DARK, [
            "MediaPipe Pose Kinematics: 33 3D skeletal landmarks",
            "Calculates joint angles, cadence & bilateral symmetry %",
            "Multi-Metric Scoring Engine: Power, Speed, VO2 max, Flexibility"
        ]),
        ("Step 4: Benchmarking", GOLD, [
            "Composite Talent Index (0-100) & Grade (A+, A, B, C)",
            "Sport Suitability Matrix (Weightlifting, Boxing, Football, Archery)",
            "AI Head-to-Head Comparative Scouting vs National Stars"
        ]),
        ("Step 5: Coach & NCOE", EMERALD, [
            "Automated Coach Dossier & Breakthrough alert",
            "SAI/NIS Qualified Coach verification & trial invitation",
            "Direct feeder into Khelo India (KITI) & TOPS developmental pool"
        ])
    ]

    for i, (title, color, items) in enumerate(steps_data):
        x_pos = Inches(0.8) + i * (step_w + gap)
        sc = add_card(s2, x_pos, y_step, step_w, step_h, bg_color=WHITE, border_color=color)
        tfs = sc.text_frame
        tfs.word_wrap = True
        tfs.margin_left = tfs.margin_right = Inches(0.18)
        tfs.margin_top = Inches(0.2)

        # Step Pill/Number
        p_num = tfs.paragraphs[0]
        p_num.text = f"PHASE {i+1}"
        p_num.font.size = Pt(10)
        p_num.font.bold = True
        p_num.font.color.rgb = color

        p_t = tfs.add_paragraph()
        p_t.text = title.split(": ")[1]
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY_DARK
        p_t.space_before = Pt(3)

        for item in items:
            p_i = tfs.add_paragraph()
            p_i.text = "• " + item
            p_i.font.size = Pt(10)
            p_i.font.color.rgb = TEXT_MUTED
            p_i.space_before = Pt(6)

    # Bottom Architecture Summary Strip
    bot_strip = add_card(s2, Inches(0.8), Inches(6.1), Inches(11.73), Inches(0.9), bg_color=NAVY_DARK, border_color=NAVY_DARK)
    tf_bs = bot_strip.text_frame
    tf_bs.word_wrap = True
    tf_bs.margin_left = tf_bs.margin_right = Inches(0.3)
    tf_bs.margin_top = Inches(0.15)
    p_b1 = tf_bs.paragraphs[0]
    p_b1.text = "CORE ARCHITECTURE: FastAPI Backend + MediaPipe Computer Vision + Google Gemini AI Coach + SQLite/PostgreSQL + React 19 Frontend"
    p_b1.font.size = Pt(11)
    p_b1.font.bold = True
    p_b1.font.color.rgb = GOLD
    p_b2 = tf_bs.add_paragraph()
    p_b2.text = "Fully decoupled micro-services architecture engineered for offline mobile caching in low-bandwidth hill areas and scalable government cloud deployment."
    p_b2.font.size = Pt(10)
    p_b2.font.color.rgb = RGBColor(226, 232, 240)


    # =========================================================================
    # SLIDE 3: Key Features, AI Methodology & Technical Feasibility
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, LIGHT_BG)
    add_header(s3, "AI Methodology, Core Features & Feasibility", "TECHNICAL DEEP DIVE")

    feat_w = Inches(5.75)
    feat_h = Inches(2.55)

    features = [
        ("1. Standard Battery of Tests (SAI Form K-01)", SAFFRON, Inches(0.8), Inches(1.6), [
            "Physical Measurements: Height, Body Weight, Wingspan-to-Height Ape Index.",
            "Explosive Power: Sargent Vertical Jump test (measured against age norms).",
            "Speed & Acceleration: Standard 30m sprint test with automated split analysis.",
            "Cardiovascular Endurance: Multi-stage 20m Shuttle Run (Beep Test) for VO2 max estimation.",
            "Flexibility: Standard Sit-and-Reach test for hamstring & lower spine mobility."
        ]),
        ("2. Computer Vision Biomechanics & Kinematics", BLUE_ACCENT, Inches(6.78), Inches(1.6), [
            "MediaPipe Pose: Real-time 33-point skeletal landmark detection from smartphone camera.",
            "Joint Angular Kinematics: Real-time angle calculation across knees, hips, and shoulders.",
            "Bilateral Symmetry Scoring: Quantifies left-vs-right asymmetry (%) to detect imbalances.",
            "Movement Quality & Rep Cadence: Form score (0-100) ensuring verified exercise execution."
        ]),
        ("3. Composite Talent Index & Recommendation", NAVY_DARK, Inches(0.8), Inches(4.35), [
            "Composite Talent Index Formula: Power (25%) + Speed (20%) + Endurance (20%) + Flexibility (15%) + Technique (20%).",
            "Breakthrough Probability Engine: Identifies grassroots prodigies outperforming peer cohorts.",
            "Sport Suitability Algorithm: Maps biometric traits to high-potential Olympic sports (Weightlifting, Boxing, Football, Archery, Wushu)."
        ]),
        ("4. National Khel ID & Verified Coach Pipeline", EMERALD, Inches(6.78), Inches(4.35), [
            "Single Source of Truth: Standardized Khel ID (KHEL-IND-MN-YYYY-XXXX) linked to NSIS.",
            "Dronacharya / NIS Coach Verification: Qualified coaches review video, metrics, and trial history.",
            "Automated Athlete Dossiers: One-click printable technical dossiers for scouting committees.",
            "Head-to-Head Comparative Scouting: Direct physical delta benchmarking against national stars."
        ])
    ]

    for title, color, x, y, bullets in features:
        fc = add_card(s3, x, y, feat_w, feat_h, bg_color=WHITE, border_color=color)
        tff = fc.text_frame
        tff.word_wrap = True
        tff.margin_left = tff.margin_right = Inches(0.25)
        tff.margin_top = Inches(0.2)

        p = tff.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color

        for b in bullets:
            p_b = tff.add_paragraph()
            p_b.text = "• " + b
            p_b.font.size = Pt(10.5)
            p_b.font.color.rgb = TEXT_DARK
            p_b.space_before = Pt(4)


    # =========================================================================
    # SLIDE 4: Competitive Advantage & Uniqueness
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, LIGHT_BG)
    add_header(s4, "Competitive Advantage & Distinctive Uniqueness", "MARKET BENCHMARKING & UNIQUENESS")

    # Comparison Table
    table_shape = s4.shapes.add_table(5, 4, Inches(0.8), Inches(1.6), Inches(11.73), Inches(2.7))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(2.8)
    tbl.columns[1].width = Inches(2.9)
    tbl.columns[2].width = Inches(3.0)
    tbl.columns[3].width = Inches(3.03)

    headers = ["Evaluation Metric", "Conventional Scouting", "Commercial Solutions", "AI4Seva Portal"]
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY_DARK if j < 3 else EMERALD
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

    table_data = [
        ("Cost per Athlete Assessed", "₹2,500 - ₹5,000 (Travel, lodging, camp logistics)", "₹15,000+ (Force plates, GPS vests, wearables)", "Near Zero (Existing smartphone camera & school field)"),
        ("Geographical Reach", "Confined to urban centers & state capitals", "Restricted to elite clubs & private academies", "Pan-India grassroots reach (remote hill & tribal villages)"),
        ("Scouting Objectivity & Bias", "Subjective coach impressions; prone to nepotism", "Data-rich but single-sport isolated metrics", "100% objective, normalized AI percentiles & Khel ID"),
        ("Talent Pathway Integration", "Disconnected paper records; lost across federations", "Private proprietary data locked in walled gardens", "Direct pipeline to SAI NCOEs, Khelo India (KITI) & TOPS")
    ]

    for i, row in enumerate(table_data):
        for j, val in enumerate(row):
            cell = tbl.cell(i+1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if j < 3 else RGBColor(236, 253, 245)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10)
            p.font.bold = (j == 0 or j == 3)
            p.font.color.rgb = NAVY_DARK if j != 3 else RGBColor(6, 95, 70)

    # 3 Distinctive Uniqueness Callout Cards below table
    u_w = Inches(3.75)
    u_h = Inches(2.4)
    y_u = Inches(4.55)

    unique_pillars = [
        ("1. Head-to-Head Comparative Scouting", SAFFRON, [
            "Direct benchmark delta of a 16-yr-old prospect against national champions (e.g. Paominthang vs Suresh Wangjam).",
            "Quantifies speed, vertical leap, and VO2 max gap with AI scout commentary."
        ]),
        ("2. Multi-Sport Cross Suitability", BLUE_ACCENT, [
            "Prevents athlete misdirection by testing general athletic foundation.",
            "Automatically redirects an athlete with high explosive power from football to Olympic weightlifting or boxing."
        ]),
        ("3. Anti-Fraud & Verified SAI Pipeline", EMERALD, [
            "Official Khel ID prevents age manipulation and duplicate registration.",
            "Two-stage verification with NIS coaches ensures transparent, tamper-proof merit."
        ])
    ]

    for i, (title, color, bullets) in enumerate(unique_pillars):
        x_pos = Inches(0.8) + i * Inches(3.99)
        uc = add_card(s4, x_pos, y_u, u_w, u_h, bg_color=WHITE, border_color=color)
        tfu = uc.text_frame
        tfu.word_wrap = True
        tfu.margin_left = tfu.margin_right = Inches(0.2)
        tfu.margin_top = Inches(0.2)

        p = tfu.paragraphs[0]
        p.text = title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = color

        for b in bullets:
            p_b = tfu.add_paragraph()
            p_b.text = "• " + b
            p_b.font.size = Pt(10)
            p_b.font.color.rgb = TEXT_MUTED
            p_b.space_before = Pt(4)


    # =========================================================================
    # SLIDE 5: Business Model, Scalability & Financial Sustainability
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, LIGHT_BG)
    add_header(s5, "Business Model, Scalability & Financial Sustainability", "SUSTAINABILITY & SCALABILITY")

    col_w = Inches(3.75)
    col_h = Inches(5.3)
    y_col = Inches(1.6)

    # Column 1: Business Model & Sustainability Beyond Hackathon
    col1 = add_card(s5, Inches(0.8), y_col, col_w, col_h, bg_color=WHITE, border_color=BLUE_ACCENT)
    t1 = col1.text_frame
    t1.word_wrap = True
    t1.margin_left = t1.margin_right = Inches(0.22)
    t1.margin_top = Inches(0.25)
    p = t1.paragraphs[0]
    p.text = "Business Model & Sustainability"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    points1 = [
        ("Institutional G2G / G2C SaaS", "Direct operational integration with Ministry of Youth Affairs & Sports (MYAS), Sports Authority of India (SAI), and Khelo India / NSTSS."),
        ("Public-Private-Partnership (PPP)", "Collaborate with State Sports Directorates and national sports federations for annual platform maintenance & verification logistics."),
        ("CSR Sports Foundation Alliances", "Funding tie-ups with Reliance Foundation Youth Sports, JSW Sports, and Tata Trusts Olympic Training programs."),
        ("Open Standard Maintenance", "Lightweight Python/FastAPI architecture with low server compute overhead, maintained by university & sports institute developers.")
    ]
    for sub, desc in points1:
        p_sub = t1.add_paragraph()
        p_sub.text = sub
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = NAVY_DARK
        p_sub.space_before = Pt(8)
        p_desc = t1.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = TEXT_MUTED

    # Column 2: Revenue Generation & Financial Model
    col2 = add_card(s5, Inches(4.79), y_col, col_w, col_h, bg_color=WHITE, border_color=EMERALD)
    t2 = col2.text_frame
    t2.word_wrap = True
    t2.margin_left = t2.margin_right = Inches(0.22)
    t2.margin_top = Inches(0.25)
    p = t2.paragraphs[0]
    p.text = "Revenue Generation Model"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    points2 = [
        ("B2B Private Academy Licensing", "Tiered subscription for private academies, school sports leagues, ISL & IPL development hubs for advanced kinematics & scouting."),
        ("Tournament Scouting As-a-Service", "Event-based licensing fee for district and state tournament organizers to automate talent identification and digital certification."),
        ("Talent Discovery Intelligence API", "API access for national leagues and professional franchises seeking verified physical benchmark data on grassroots draft prospects."),
        ("Long-Term Cost Efficiency", "Reduces SAI talent scout trial costs by over 80% while expanding scouting throughput 20-fold.")
    ]
    for sub, desc in points2:
        p_sub = t2.add_paragraph()
        p_sub.text = sub
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = NAVY_DARK
        p_sub.space_before = Pt(8)
        p_desc = t2.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = TEXT_MUTED

    # Column 3: Market Potential & Scalability
    col3 = add_card(s5, Inches(8.78), y_col, col_w, col_h, bg_color=WHITE, border_color=SAFFRON)
    t3 = col3.text_frame
    t3.word_wrap = True
    t3.margin_left = t3.margin_right = Inches(0.22)
    t3.margin_top = Inches(0.25)
    p = t3.paragraphs[0]
    p.text = "Market Potential & Scalability"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = SAFFRON

    points3 = [
        ("Massive Addressable Base", "Over 100M+ youth athletes in schools, colleges, and rural sports clubs across 700+ districts in India."),
        ("Pan-India Multi-State Expansion", "Proven pilot in Manipur sports powerhouse expands readily to all 28 States and 8 Union Territories."),
        ("Cross-Department Utility", "Directly adaptable for Indian Armed Forces, Central Armed Police Forces (CRPF, BSF), and State Police physical recruitment exams."),
        ("Pathway to Olympic Podiums", "Creates India's unified grassroots sports registry feeding athletes into TOPS for Los Angeles 2028 and Brisbane 2032.")
    ]
    for sub, desc in points3:
        p_sub = t3.add_paragraph()
        p_sub.text = sub
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = NAVY_DARK
        p_sub.space_before = Pt(8)
        p_desc = t3.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = TEXT_MUTED

    # Save presentation
    output_path = os.path.join(r"c:\Users\Aryan\OneDrive\Desktop\AI4Seva", "AI4Seva_YAS01_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()
